"""Independent oracle cross-check of the reader fixtures.

Reads the hand-written/reader fixtures with openpyxl (never with gridxcel) and
prints what the oracle sees, so the reader's assertions can be diffed.
"""
import os, sys, zipfile, datetime
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

FIX = os.path.join("gridxcel", "test", "fixtures")

for name in ["reader-handmade", "reader-1904", "reader-comments", "reader-stored-big"]:
    path = os.path.join(FIX, name + ".xlsx")
    print("=====", name)
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as exc:
        print("   ORACLE REFUSED:", type(exc).__name__, exc)
        continue
    print("   sheets:", wb.sheetnames, "epoch:", wb.epoch)
    for ws in wb.worksheets:
        print("   --", ws.title, "dims", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column,
              "freeze", ws.freeze_panes, "autofilter", ws.auto_filter.ref,
              "merges", sorted(str(m) for m in ws.merged_cells.ranges))
        if ws.title == "Stored":
            print("      A1", ws["A1"].value, "B1", ws["B1"].value, "C6000", ws["C6000"].value)
            continue
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12)):
            out = []
            for cell in row:
                if cell.value is None:
                    continue
                v = cell.value
                if isinstance(v, (datetime.datetime, datetime.date)):
                    v = v.isoformat()
                out.append("%s=%r(%s,%s)" % (cell.coordinate, v, cell.data_type, cell.number_format))
            if out:
                print("      ", "; ".join(out))
