"""Generate reader fixtures whose exact content is declared by hand.

Everything here is written with the standard library only (zipfile) or with
openpyxl, never with gridxcel, so it stays an independent oracle.
"""
import os, sys, zipfile, shutil
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "test", "fixtures")
os.makedirs(OUT, exist_ok=True)

NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
RNS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
XMLDECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

CONTENT_TYPES = XMLDECL + (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
    '{sheetOverrides}'
    '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
    '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
    '</Types>')

PKG_RELS = XMLDECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
    '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
    '</Relationships>')

CORE = XMLDECL + (
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
    ' xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/">'
    '<dc:title>Handmade reader fixture</dc:title><dc:creator>reader-eng</dc:creator>'
    '<cp:lastModifiedBy>reader-eng</cp:lastModifiedBy><cp:keywords>one, two</cp:keywords>'
    '<dc:description>declared by hand</dc:description><cp:category>tests</cp:category>'
    '<dcterms:created>2024-01-01T00:00:00Z</dcterms:created>'
    '<dcterms:modified>2024-06-01T12:00:00Z</dcterms:modified></cp:coreProperties>')

APP = XMLDECL + ('<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                 '<Application>gridxcel reader fixtures</Application><Company>none</Company></Properties>')

STYLES = XMLDECL + (
    '<styleSheet %s><numFmts count="3">'
    '<numFmt numFmtId="164" formatCode="yyyy-mm-dd"/>'
    '<numFmt numFmtId="165" formatCode="0.00%%"/>'
    '<numFmt numFmtId="166" formatCode="[$-409]d\\-mmm\\-yy;@"/>'
    '</numFmts>'
    '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
    '<font><b val="1"/><sz val="11"/><name val="Calibri"/><color rgb="FFFF0000"/></font></fonts>'
    '<fills count="2"><fill><patternFill patternType="none"/></fill>'
    '<fill><patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor indexed="64"/></patternFill></fill></fills>'
    '<borders count="1"><border><left style="thin"><color rgb="FF000000"/></left><right/><top/><bottom/><diagonal/></border></borders>'
    '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    '<cellXfs count="5">'
    '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    '<xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
    '<xf numFmtId="165" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
    '<xf numFmtId="0" fontId="1" fillId="1" borderId="0" xfId="0" applyFont="1" applyFill="1">'
    '<alignment horizontal="center" wrapText="1"/></xf>'
    '<xf numFmtId="166" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
    '</cellXfs>'
    '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
    '</styleSheet>') % NS


def sheet_override(n):
    return ('<Override PartName="/xl/worksheets/sheet%d.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' % n)


def build(parts, path):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts:
            z.writestr(name, data.encode("utf-8"))
    print("wrote", path, os.path.getsize(path), "bytes")


# ---------------------------------------------------------------------------
# 1. handmade.xlsx: cell types, formulas, offsets, attributes, 1904 off
# ---------------------------------------------------------------------------
SHARED = XMLDECL + ('<sst %s count="4" uniqueCount="4">'
                    '<si><t>plain</t></si>'
                    '<si><r><rPr><b val="1"/><color rgb="FFFF0000"/></rPr><t>Bold</t></r>'
                    '<r><t xml:space="preserve"> red</t></r><r><t> plain</t></r></si>'
                    '<si><t xml:space="preserve">  lead and trail  </t></si>'
                    '<si><t>%s</t></si>'
                    '</sst>') % (NS, "&lt;escaped&gt; &amp; \u00e9\u4e2d")

TYPES2 = XMLDECL + (
    '<worksheet %s %s><dimension ref="A1:D11"/>'
    '<sheetPr><tabColor rgb="FF00B050"/></sheetPr>'
    '<sheetViews><sheetView tabSelected="1" workbookViewId="0">'
    '<pane xSplit="2" ySplit="1" topLeftCell="C2" state="frozen"/></sheetView></sheetViews>'
    '<sheetFormatPr defaultRowHeight="15"/>'
    '<cols><col min="1" max="1" width="22" customWidth="1" style="3" outlineLevel="1"/>'
    '<col min="2" max="2" hidden="1" width="9" outlineLevel="2"/></cols>'
    '<sheetData>'
    '<row r="1" ht="22.5" customHeight="1" s="3" customFormat="1" outlineLevel="1">'
    '<c r="A1" t="d"><v>2024-01-01T00:00:00</v></c>'
    '<c r="B1"><f t="shared" ref="B1:B3" si="0">A1*2</f><v>7</v></c>'
    '<c r="C1"><f t="array" ref="C1:C3">SUM(A1:A3)</f><v>0</v></c>'
    '<c r="D1"><f t="shared" ref="D1:D2" si="1">$A$1+A1</f><v>0</v></c></row>'
    '<row r="2"><c r="A2" t="e"><v>#DIV/0!</v></c>'
    '<c r="B2"><f t="shared" si="0"/><v>0</v></c>'
    '<c r="D2"><f t="shared" si="1"/><v>0</v></c></row>'
    '<row r="3" hidden="1"><c r="A3" t="str"><f>CONCATENATE("a","b")</f><v>ab</v></c>'
    '<c r="B3"><f t="shared" si="0"/><v>0</v></c></row>'
    '<row r="4"><c r="A4" t="s"><v>1</v></c><c r="B4" t="s"><v>2</v></c>'
    '<c r="C4" t="s"><v>3</v></c><c r="D4" t="s"><v>0</v></c></row>'
    '<row r="5"><c r="A5" t="inlineStr"><is><r><t>Hello </t></r><r><t>world</t></r></is></c>'
    '<c r="B5" t="inlineStr"><is><t xml:space="preserve">  spaced  </t></is></c>'
    '<c r="C5" t="inlineStr"><is><t/></is></c></row>'
    '<row r="6"><c r="A6" t="b"><v>0</v></c><c r="B6" t="b"><v>1</v></c>'
    '<c r="C6" t="n"><v>1e3</v></c><c r="D6"><v>-0.5</v></c></row>'
    '<row r="7"><c r="A7" s="1"><v>45000</v></c><c r="B7" s="2"><v>0.125</v></c>'
    '<c r="C7" s="3"><v>42</v></c><c r="D7" s="4"><v>45000</v></c></row>'
    '<row r="8"><c><v>1</v></c><c><v>2</v></c><c><v>3</v></c></row>'
    '<row r="9"><c r="A9" s="3"/><c r="B9"/></row>'
    '<row r="10"><c r="A10" t="d"><v>2024-03-15T14:30:45</v></c>'
    '<c r="B10" t="n"><v>2</v></c></row>'
    '<row r="11"><c r="A11"><f>TODAY()</f><v>45000</v></c>'
    '<c r="B11" s="1"><f>TODAY()</f><v>45000</v></c>'
    '<c r="C11" s="1"><v>45000</v></c></row>'
    '</sheetData>'
    '<autoFilter ref="A1:D4"/>'
    '<mergeCells count="2"><mergeCell ref="A12:C13"/><mergeCell ref="D4:D5"/></mergeCells>'
    '<hyperlinks><hyperlink ref="A4" r:id="rId1" tooltip="external tip"/>'
    '<hyperlink ref="B4" location="Types2!C6" display="jump"/></hyperlinks>'
    '</worksheet>') % (NS, RNS)

TYPES2_RELS = XMLDECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"'
    ' Target="https://example.org/hand" TargetMode="External"/></Relationships>')

OFFSET = XMLDECL + (
    '<worksheet %s><dimension ref="C5:E6"/><sheetData>'
    '<row r="5"><c r="C5"><v>1</v></c><c r="E5"><v>2</v></c></row>'
    '<row r="6"><c r="D6"><v>3</v></c></row>'
    '</sheetData></worksheet>') % NS

NODATA = XMLDECL + ('<worksheet %s><dimension ref="A1"/></worksheet>') % NS

HIDDEN = XMLDECL + ('<worksheet %s><sheetData><row r="1"><c r="A1"><v>9</v></c></row></sheetData></worksheet>') % NS

WORKBOOK = XMLDECL + (
    '<workbook %s %s><workbookPr/><bookViews><workbookView/></bookViews><sheets>'
    '<sheet name="Types2" sheetId="1" r:id="rId1"/>'
    '<sheet name="Offset" sheetId="2" r:id="rId2"/>'
    '<sheet name="NoData" sheetId="3" r:id="rId3"/>'
    '<sheet name="Hidden" sheetId="4" state="hidden" r:id="rId4"/>'
    '</sheets><definedNames>'
    '<definedName name="MyRange">Types2!$A$1:$A$10</definedName>'
    '<definedName name="LocalName" localSheetId="1" hidden="1">Offset!$C$5</definedName>'
    '</definedNames><calcPr calcId="1"/></workbook>') % (NS, RNS)

WORKBOOK_RELS = XMLDECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>'
    '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet3.xml"/>'
    '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet4.xml"/>'
    '<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    '<Relationship Id="rId6" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
    '</Relationships>')

CT = CONTENT_TYPES.replace("{sheetOverrides}", "".join(sheet_override(n) for n in (1, 2, 3, 4)))
build([
    ("[Content_Types].xml", CT),
    ("_rels/.rels", PKG_RELS),
    ("docProps/core.xml", CORE),
    ("docProps/app.xml", APP),
    ("xl/workbook.xml", WORKBOOK),
    ("xl/_rels/workbook.xml.rels", WORKBOOK_RELS),
    ("xl/styles.xml", STYLES),
    ("xl/sharedStrings.xml", SHARED),
    ("xl/worksheets/sheet1.xml", TYPES2),
    ("xl/worksheets/_rels/sheet1.xml.rels", TYPES2_RELS),
    ("xl/worksheets/sheet2.xml", OFFSET),
    ("xl/worksheets/sheet3.xml", NODATA),
    ("xl/worksheets/sheet4.xml", HIDDEN),
], os.path.join(OUT, "reader-handmade.xlsx"))

# 2. 1904 date system, one date cell with a date format.
D1904_SHEET = XMLDECL + (
    '<worksheet %s><dimension ref="A1:B2"/><sheetData>'
    '<row r="1"><c r="A1" s="1" t="n"><v>43830.5</v></c>'
    '<c r="B1" t="d"><v>2024-01-01T12:00:00</v></c></row>'
    '<row r="2"><c r="A2"><v>1</v></c><c r="B2" s="1" t="n"><v>1</v></c></row>'
    '</sheetData></worksheet>') % NS
WB1904 = XMLDECL + (
    '<workbook %s %s><workbookPr date1904="1"/><sheets>'
    '<sheet name="D1904" sheetId="1" r:id="rId1"/></sheets></workbook>') % (NS, RNS)
WB1904_RELS = XMLDECL + (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
    '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    '</Relationships>')
CT1904 = CONTENT_TYPES.replace("{sheetOverrides}", sheet_override(1)).replace(
    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>', "")
build([
    ("[Content_Types].xml", CT1904),
    ("_rels/.rels", PKG_RELS),
    ("docProps/core.xml", CORE),
    ("docProps/app.xml", APP),
    ("xl/workbook.xml", WB1904),
    ("xl/_rels/workbook.xml.rels", WB1904_RELS),
    ("xl/styles.xml", STYLES),
    ("xl/worksheets/sheet1.xml", D1904_SHEET),
], os.path.join(OUT, "reader-1904.xlsx"))

# 3. A valid zip whose worksheet part is cut mid-element.
src = open(os.path.join(OUT, "openpyxl-written.xlsx"), "rb").read()
with zipfile.ZipFile(os.path.join(OUT, "openpyxl-written.xlsx")) as zin:
    sheet1 = zin.read("xl/worksheets/sheet1.xml").decode("utf-8")
    others = [(i.filename, zin.read(i.filename)) for i in zin.infolist() if i.filename != "xl/worksheets/sheet1.xml"]
cut = sheet1[: int(len(sheet1) * 0.55)]
with zipfile.ZipFile(os.path.join(OUT, "reader-cut-sheet.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
    for name, data in others:
        z.writestr(name, data)
    z.writestr("xl/worksheets/sheet1.xml", cut.encode("utf-8"))
print("wrote reader-cut-sheet.xlsx", os.path.getsize(os.path.join(OUT, "reader-cut-sheet.xlsx")), "bytes; sheet1 cut at", len(cut), "of", len(sheet1))

# 4. Comments written by openpyxl (oracle), plus a comment part name check.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "WithComments"
ws["A1"] = "has comment"
ws["A1"].comment = Comment("first note", "reader-eng")
ws["C2"] = "also"
ws["C2"].comment = Comment("second note", "oracle")
ws2 = wb.create_sheet("NoComments")
ws2["A1"] = 1
wb.save(os.path.join(OUT, "reader-comments.xlsx"))
with zipfile.ZipFile(os.path.join(OUT, "reader-comments.xlsx")) as z:
    print("comments fixture parts:", [n for n in z.namelist() if "comment" in n or "rels" in n])
    for n in z.namelist():
        if n.endswith(".rels") and "sheet" in n:
            print("  ", n, "->", z.read(n).decode("utf-8")[:400])

# 4b. Deliberately malformed cells: every tolerant-reader warning path.
BADSHEET = XMLDECL + (
    '<worksheet %s><dimension ref="A1:E2"/><sheetData>'
    '<row r="1"><c r="A1" t="d"><v>not-a-date</v></c>'
    '<c r="B1"><v>abc</v></c>'
    '<c r="C1" t="s"><v>99</v></c>'
    '<c r="D1" t="q"><v>x</v></c>'
    '<c r="E1" t="s"><v>0</v></c></row>'
    '<c r="A2"><v>5</v></c>'
    '</sheetData></worksheet>') % NS
BADSHARED = XMLDECL + ('<sst %s count="1" uniqueCount="1"><si><t>only</t></si></sst>') % NS
CTBAD = CONTENT_TYPES.replace("{sheetOverrides}", sheet_override(1))
build([
    ("[Content_Types].xml", CTBAD),
    ("_rels/.rels", PKG_RELS),
    ("docProps/core.xml", CORE),
    ("docProps/app.xml", APP),
    ("xl/workbook.xml", XMLDECL + ('<workbook %s %s><sheets><sheet name="Bad" sheetId="1" r:id="rId1"/></sheets></workbook>' % (NS, RNS))),
    ("xl/_rels/workbook.xml.rels", XMLDECL + ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
        '</Relationships>')),
    ("xl/styles.xml", STYLES),
    ("xl/sharedStrings.xml", BADSHARED),
    ("xl/worksheets/sheet1.xml", BADSHEET),
], os.path.join(OUT, "reader-badcells.xlsx"))

# 4c. A workbook that declares no sheets at all.
build([
    ("[Content_Types].xml", CONTENT_TYPES.replace("{sheetOverrides}", "")),
    ("_rels/.rels", PKG_RELS),
    ("docProps/core.xml", CORE),
    ("xl/workbook.xml", XMLDECL + ('<workbook %s><sheets/></workbook>' % NS)),
    ("xl/_rels/workbook.xml.rels", XMLDECL + ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')),
], os.path.join(OUT, "reader-nosheets.xlsx"))

# 4d. pandas workbook with a sharedStrings part cut in half (readable zip,
# unreadable table), which is the "unreadable sharedStrings" warning path.
with zipfile.ZipFile(os.path.join(OUT, "pandas.xlsx")) as zin:
    sst = zin.read("xl/sharedStrings.xml")
    sst_cut = sst[: int(len(sst) * 0.5)]
    with zipfile.ZipFile(os.path.join(OUT, "reader-cut-sst.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
        for it in zin.infolist():
            z.writestr(it.filename, sst_cut if it.filename == "xl/sharedStrings.xml" else zin.read(it.filename))
print("wrote reader-nosheets.xlsx and reader-cut-sst.xlsx (sst cut to", len(sst_cut), "of", len(sst), "bytes)")

# 5. A STORED sheet big enough to need several scanner slices (~600 KB).
rows = []
for i in range(1, 6001):
    rows.append('<row r="%d"><c r="A%d" t="inlineStr"><is><t>r%d</t></is></c>'
                '<c r="B%d"><v>%d</v></c><c r="C%d"><v>%s</v></c></row>'
                % (i, i, i, i, i, i, i * 1.5))
stored_sheet = XMLDECL + ('<worksheet %s><dimension ref="A1:C6000"/><sheetData>%s</sheetData></worksheet>'
                          % (NS, "".join(rows)))
wb5 = XMLDECL + ('<workbook %s %s><sheets><sheet name="Stored" sheetId="1" r:id="rId1"/></sheets></workbook>' % (NS, RNS))
wb5_rels = XMLDECL + ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                      '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                      '</Relationships>')
CT5 = CONTENT_TYPES.replace("{sheetOverrides}", sheet_override(1)).replace(
    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>', "")
with zipfile.ZipFile(os.path.join(OUT, "reader-stored-big.xlsx"), "w", zipfile.ZIP_STORED) as z:
    z.writestr("[Content_Types].xml", CT5)
    z.writestr("_rels/.rels", PKG_RELS)
    z.writestr("docProps/core.xml", CORE)
    z.writestr("xl/workbook.xml", wb5)
    z.writestr("xl/_rels/workbook.xml.rels", wb5_rels)
    z.writestr("xl/worksheets/sheet1.xml", stored_sheet.encode("utf-8"))
print("wrote reader-stored-big.xlsx", os.path.getsize(os.path.join(OUT, "reader-stored-big.xlsx")), "bytes; sheet xml", len(stored_sheet))
