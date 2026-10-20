
"""Generate test/fixtures/golden.xlsx plus test/fixtures/golden.json.

The manifest describes exactly what an independent implementation (openpyxl)
sees in the file, so any reader can be diffed against it without re-running
Python. Regenerate with:  python test/make_golden.py
"""
import json, os, datetime, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

here = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(here, "fixtures")
os.makedirs(out, exist_ok=True)
path = os.path.join(out, "golden.xlsx")

wb = openpyxl.Workbook()

# --- Sheet 1: types, formatting, formulas -----------------------------------
ws = wb.active
ws.title = "Types"
ws["A1"] = "text"; ws["A1"].font = Font(bold=True)
ws["B1"] = 42
ws["C1"] = True
ws["D1"] = datetime.datetime(2024, 3, 15, 14, 30, 45)
ws["D1"].number_format = "yyyy-mm-dd hh:mm:ss"
ws["E1"] = 0.125
ws["E1"].number_format = "0.00%"
ws["F1"] = 1234.5
ws["F1"].number_format = "#,##0.00"
ws["G1"] = "line1\nline2"
ws["H1"] = "quote\"and,comma"
ws["I1"] = "\u65e5\u672c\u8a9e \U0001f600 \u00e9\u00e8"
ws["J1"] = 45292
ws["J1"].number_format = "yyyy-mm-dd"
ws["A2"] = "=B1*2"
ws["B2"] = "=SUM(B1:B1)"
ws["C2"] = -0.0001234
ws["C2"].number_format = "0.00E+00"
ws["D2"] = None
ws["E2"] = 0
ws["F2"] = -5
ws["G2"] = True
ws["H2"] = 3.14159265358979
ws["I2"] = "  padded  "
ws["J2"] = 1e15

# --- Sheet 2: layout --------------------------------------------------------
ws2 = wb.create_sheet("Layout")
ws2["A1"] = "merged"
ws2.merge_cells("A1:C1")
ws2.merge_cells("A3:B4")
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 12.5
ws2.row_dimensions[1].height = 28
ws2.row_dimensions[5].hidden = True
ws2.column_dimensions["D"].hidden = True
ws2.freeze_panes = "B2"
ws2["A7"] = "after hidden"
ws2["C7"] = 7
ws2["A8"] = "tab colour target"
ws2.sheet_properties.tabColor = "FF0000"

# --- Sheet 3: volume + sparse columns ---------------------------------------
ws3 = wb.create_sheet("Volume")
for r in range(1, 201):
    ws3.cell(row=r, column=1, value="row%d" % r)
    ws3.cell(row=r, column=2, value=r * 1.5)
    ws3.cell(row=r, column=26, value="far%d" % r)   # column Z, sparse gap
for r in range(1, 201):
    ws3.cell(row=r, column=4, value=r % 7)

# --- Sheet 4: empty ---------------------------------------------------------
wb.create_sheet("Empty")

wb.properties.title = "gridxcel golden fixture"
wb.properties.creator = "make_golden.py"
wb.save(path)

# --- Manifest ---------------------------------------------------------------
types = openpyxl.load_workbook(path, data_only=False)
data = openpyxl.load_workbook(path, data_only=True)

def cell_json(c, value_only=False):
    if c.value is None:
        return None
    v = c.value
    if isinstance(v, (datetime.datetime, datetime.date)):
        v = v.isoformat()
    elif isinstance(v, bool):
        v = v
    elif isinstance(v, float):
        v = float(v)
    t = c.data_type
    return {"v": v, "t": t, "nf": c.number_format, "addr": c.coordinate}

manifest = {
    "file": "golden.xlsx",
    "sheets": types.sheetnames,
    "cellDates": {},
    "formulas": {},
    "merges": {ws.title: sorted(str(m) for m in ws.merged_cells.ranges) for ws in types.worksheets},
    "dimensions": {ws.title: ws.dimensions for ws in types.worksheets},
    "volume": {"rows": 200, "colA": "row1", "colB1": 1.5, "colZ1": "far1", "colD1": 1},
}
for name in ["Types", "Layout"]:
    s = types[name]
    manifest["cellDates"][name] = {}
    for row in s.iter_rows():
        for c in row:
            if c.value is not None:
                manifest["cellDates"][name][c.coordinate] = cell_json(c)
    d = data[name]
    manifest["formulas"][name] = {c.coordinate: c.value for row in d.iter_rows() for c in row if c.value is not None and c.data_type == "f"}

# expected computed values for the formulas (data_only needs a cached value,
# which openpyxl cannot compute - state them explicitly instead)
manifest["expectedFormulaResults"] = {"Types!A2": 84, "Types!B2": 42}

with open(os.path.join(here, "fixtures", "golden.json"), "w", encoding="utf-8") as fh:
    json.dump(manifest, fh, indent=1, ensure_ascii=False)

print("wrote", path, os.path.getsize(path), "bytes")
print("sheets:", manifest["sheets"])
print("merges:", manifest["merges"])
print("dimensions:", manifest["dimensions"])
print("Types!A1:", json.dumps(manifest["cellDates"]["Types"]["A1"], ensure_ascii=False))
print("Types!D1:", json.dumps(manifest["cellDates"]["Types"]["D1"], ensure_ascii=False))
print("Types!J1:", json.dumps(manifest["cellDates"]["Types"]["J1"], ensure_ascii=False))
print("Types!I1:", json.dumps(manifest["cellDates"]["Types"]["I1"], ensure_ascii=False))
print("formulas:", json.dumps(manifest["formulas"], ensure_ascii=False))
